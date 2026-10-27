"""
発表のスライド（10 枚、7 分、理系の大学生向け）を作る。構成は ../発表の構成.md。前回の中間発表のファイルを土台にして、
表紙（ロゴ・題名・所属）とスライドの大きさを引き継ぎ、2 枚目以降を作る。

2 枚目以降の作りは、参考にした調査報告書（../report_202607241552_0.pdf）に合わせる：2 行の結論の文のタイトル、
色の付いた小見出し、右上の発表の流れ、左に濃い灰色の項目の欄を置いた表、角の立った四角、■・➢ の箇条書き、
下向きの台形と濃い灰色の帯のまとめ、*1 の脚注、左下のページ番号。文は「である」調か体言止め。

数値はすべて実験の記録から（全体の記録の第 1・32・33・90 章、総合報告書 3 の第 VI 部、
experiments/G_proof/results/summary.json、experiments/G3_confirm/results/summary.json、実験 2 の図のデータ）。
7 枚目の資産の推移は experiments/G_proof/results/*/real.csv（実験 G の日々の系列）から計算する。

    python 発表.py            （python-pptx が要る）

出力：../2026_情報工学工房発表資料.pptx（土台のファイルは変えない）
"""
import csv
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
BASE = HERE.parent.parent / '06_中間発表と途中のメモ' / '2026_情報工学工房中間発表資料_修正版claude.pptx'
OUT = HERE.parent / '2026_情報工学工房発表資料.pptx'
ROOT = HERE.parents[2]

FONT = 'メイリオ'
ACCENT = '1973EB'     # the title slide's blue
BLACK = '000000'
DGRAY = '52555A'      # label blocks and the closing band
MGRAY = '7E7F7E'      # small print
GRAY = '96999B'
LINE = 'BABBBB'
SILVER = 'D8D9D8'
LGRAY = 'F1F2F1'      # table cells
WHITE = 'FFFFFF'
RED = 'C0281E'        # ✗ only
LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE

# the steps of the talk, for the map at the top right
FLOW = ['元論文', '決算文', '大型株', '小型株', '検証']

# 実験 2 の図（資料集/02_実験/実験2_決算発表文/作成スクリプト/figs.py の fig_quintile と fig_car）と同じ計算の値。
# 点数の 5 分位ごとの、発表をはさむ 2 取引日の超過リターンの平均（%）
QUINTILE = [-1.773, -0.198, -0.250, 0.456, 1.164]
# 点数の上位・中位・下位 20% の、発表の 5 日前から 21 日後までの累積の超過リターン（%）
DAYS = list(range(-5, 22))
PATH_TOP = [0.07, 0.245, 0.306, 0.382, 0.542, 1.811, 1.808, 1.916, 2.0, 2.027, 1.992, 2.067, 2.065, 2.12, 2.119,
            2.149, 2.142, 2.104, 2.259, 2.228, 2.178, 2.238, 2.352, 2.333, 2.404, 2.382, 2.329]
PATH_MID = [0.061, 0.027, 0.107, 0.145, 0.139, 0.047, -0.087, -0.133, -0.147, -0.249, -0.307, -0.302, -0.304, -0.453,
            -0.482, -0.528, -0.525, -0.514, -0.518, -0.537, -0.437, -0.429, -0.393, -0.511, -0.454, -0.436, -0.412]
PATH_LOW = [-0.182, -0.194, -0.159, -0.185, -0.245, -2.099, -2.207, -2.24, -2.2, -2.18, -2.159, -2.261, -2.232, -2.274,
            -2.211, -2.278, -2.331, -2.283, -2.323, -2.339, -2.302, -2.298, -2.273, -2.297, -2.302, -2.241, -2.239]

# 文章を数にする 7 つの方法（第 32 章の表、第 90 章「モデルの強さの順位」、各実験の事前登録）。
# (モデル, 補足, 何を学習させたか, 補足, 学習の方法, 補足, R² の増分の下限, 上限, LLM か)
MODELS = [
    ('gpt-4o-mini', 'OpenAI、規模は非公開', '学習なし', '指示のみで −2〜+2 を回答',
     'ゼロショット', '例示なしで質問のみ', 2.01, 2.49, True),
    ('FinBERT（追加学習）', '1.1 億パラメータ', '文章 → 株価反応', '先頭 512 トークン、約 1.4 万本',
     '全体の追加学習', '全パラメータを更新（2 エポック）', 0.84, 1.01, False),
    ('ChronoGPT＋リッジ', '15.5 億、先読みなし（年次版）', '内部ベクトル → 反応', '1,536 次元、四半期ごとに再学習',
     '線形プローブ', 'LLM は固定し、リッジ回帰のみ学習', 0.28, 0.59, True),
    ('TF-IDF＋リッジ', '語彙 2 万', '単語の重み → 反応', '全文、四半期ごとに再学習',
     'リッジ回帰', '反応と結び付く単語を学習', 0.06, 0.37, False),
    ('FinBERT（学習なし）', '1.1 億パラメータ', '学習なし', '先頭 200 文を 1 文ずつ',
     '肯定 − 否定の確率', '金融文で学習済みの分類器', 0.24, 0.24, False),
    ('ChronoGPT＋LoRA', '15.5 億中 256 万（0.17%）', '文章 → 反応', '先頭 1,792 トークン、約 1.4 万本',
     'LoRA', '各層に低ランク行列を追加（7.6 時間）', 0.13, 0.13, True),
    ('辞書', 'Loughran–McDonald 金融辞書', '学習なし', '全文',
     '単語の計数', '肯定語と否定語の数の差', 0.02, 0.02, False),
]

# 実験 G の評価の期間と費用（experiments/G_proof/analyze.py の EVAL と主な費用）
EVAL = '2017-01-01'
COST = 0.003


def asset_paths(sample):
    """The small-cap BLM and its band, valued from 1 at the end of 2016 after a one-way cost of 0.3% on both sides'
    trades, at each month end; also the annual mean active return after costs (the G1/G5 value)."""
    with open(ROOT / 'experiments' / 'G_proof' / 'results' / sample / 'real.csv', encoding='utf-8') as f:
        rows = [r for r in csv.DictReader(f) if r['date'] >= EVAL]
    v_blm = v_band = 1.0
    months, blm, band, active = [], [], [], []
    for r in rows:
        v_blm *= 1 + float(r['port']) - COST * float(r['traded'])
        v_band *= 1 + float(r['bench']) - COST * float(r['traded_band'])
        active.append(float(r['active']) - COST * (float(r['traded']) - float(r['traded_band'])))
        if months and months[-1] == r['date'][:7]:
            blm[-1], band[-1] = v_blm, v_band
        else:
            months.append(r['date'][:7])
            blm.append(v_blm)
            band.append(v_band)
    # the start (end of 2016) and a year label every 12 months (the end of December is the start of the next year)
    cats = [str(2017 + i // 12) if i % 12 == 0 else '' for i in range(len(months) + 1)]
    return cats, [1.0] + blm, [1.0] + band, sum(active) / len(active) * 252


# ---------- helpers ----------

def rgb(h):
    return RGBColor.from_string(h)


def set_font(font, size=None, bold=None, color=None):
    """Size, weight, colour, and the same face for Latin and Japanese text; Japanese line breaking."""
    if size is not None:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if color is not None:
        font.color.rgb = rgb(color)
    font.name = FONT
    rPr = font._rPr
    rPr.set('lang', 'ja-JP')          # Japanese line breaking (no 。 at the start of a line)
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set('typeface', FONT)


def text(slide, x, y, w, h, paras, size=12, color=BLACK, bold=False, align=LEFT, anchor=TOP, after=0, spacing=None,
         name=None):
    """A text box. paras: list of paragraphs; a paragraph is a str or a list of (text, {size, color, bold}).
    spacing: a fixed line spacing in points (one value, or one per paragraph)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if after:
            p.space_after = Pt(after)
        if spacing is not None:
            p.line_spacing = Pt(spacing[i] if isinstance(spacing, (list, tuple)) else spacing)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            set_font(r.font, o.get('size', size), o.get('bold', bold), o.get('color', color))
    return tb


def bullets(slide, x, y, w, h, items, size=12, after=3, anchor=TOP, name=None):
    """Bullet paragraphs: ■ for level 0, ➢ for level 1. An item is (level, str or a list of (text, options))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (level, para) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else para
        for t, o in runs:
            r = p.add_run()
            r.text = t
            set_font(r.font, o.get('size', size - level), o.get('bold', False), o.get('color', BLACK))
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(Inches(0.2 + 0.22 * level)))
        pPr.set('indent', str(-Inches(0.2)))
        clr = pPr.makeelement(qn('a:buClr'), {})
        clr.append(clr.makeelement(qn('a:srgbClr'), {'val': BLACK}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn('a:buSzPct'), {'val': '75000'}))
        pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial' if level == 0 else 'Wingdings'}))
        pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '■' if level == 0 else 'Ø'}))
    return tb


def shape(slide, kind, x, y, w, h, fill, line=None, rot=0, name=None):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.0)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if rot:
        s.rotation = rot
    return s


def rect(slide, x, y, w, h, fill, line=None, name=None):
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill, line=line, name=name)


def tri(slide, x, y, w, h, fill, rot=0):
    """A small solid triangle; rot 90 points right, 180 points down."""
    return shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, x, y, w, h, fill, rot=rot)


def rule(slide, x1, y1, x2, y2, color=BLACK, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def cell(slide, x, y, w, h, paras, fill=LGRAY, size=12, color=BLACK, bold=False, align=LEFT, anchor=MIDDLE, pad=0.1,
         line=None, after=0, spacing=None, name=None):
    """A filled rectangle with text inset by pad."""
    rect(slide, x, y, w, h, fill, line=line, name=name)
    return text(slide, x + pad, y + 0.04, w - 2 * pad, h - 0.08, paras, size=size, color=color, bold=bold,
                align=align, anchor=anchor, after=after, spacing=spacing)


def label_block(slide, x, y, w, h, label, sub=None, fill=DGRAY, size=13):
    """The dark row label at the left of a table."""
    paras = [label] + ([[(sub, {'size': 10, 'bold': False})]] if sub else [])
    return cell(slide, x, y, w, h, paras, fill=fill, color=WHITE, bold=True, size=size, align=CENTER, pad=0.06)


def tag(slide, x, y, s, fill=ACCENT, d=0.24, size=10):
    """A small square with a letter or a number, at the top left of a block."""
    rect(slide, x, y, d, d, fill)
    text(slide, x, y, d, d, [s], size=size, bold=True, color=WHITE, align=CENTER, anchor=MIDDLE)


def arrow_circle(slide, cx, cy, d=0.3):
    """The black circle with » between two columns."""
    shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, BLACK)
    text(slide, cx - d / 2, cy - d / 2, d, d, ['»'], size=12, bold=True, color=WHITE, align=CENTER, anchor=MIDDLE)


def oval(slide, x, y, w, h, fill, paras, color):
    shape(slide, MSO_SHAPE.OVAL, x, y, w, h, fill)
    text(slide, x + 0.15, y, w - 0.3, h, paras, size=13, bold=True, color=color, align=CENTER, anchor=MIDDLE)


def width_of(s, size):
    """A rough width in inches of a string of Japanese and Latin characters."""
    full = sum(1 for c in s if ord(c) > 0x2000)
    return (full + 0.55 * (len(s) - full)) * size / 72


def section(slide, x, y, w, label, size=12.5):
    """A centred heading with a rule on each side: ―― label ――."""
    tw = width_of(label, size)
    cx = x + w / 2
    text(slide, cx - tw / 2 - 0.1, y, tw + 0.2, 0.32, [label], size=size, bold=True, align=CENTER, anchor=MIDDLE)
    rule(slide, x, y + 0.16, cx - tw / 2 - 0.15, y + 0.16)
    rule(slide, cx + tw / 2 + 0.15, y + 0.16, x + w, y + 0.16)


def col_heads(slide, y, cols, size=12):
    """Column headings with a rule under each; cols: (x, w, label)."""
    for x, w, label in cols:
        if label:
            text(slide, x, y, w, 0.3, [label], size=size, bold=True, align=CENTER, anchor=MIDDLE)
            rule(slide, x + 0.05, y + 0.36, x + w - 0.05, y + 0.36, width=1.0)


def funnel(slide, x, y, w, h):
    """The grey trapezoid narrowing downwards into the closing band."""
    s = shape(slide, MSO_SHAPE.TRAPEZOID, x, y, w, h, LINE, rot=180)
    s.adjustments[0] = 0.4
    return s


def closing(slide, y, h, s, size=15):
    """The dark full-width band with the slide's conclusion."""
    cell(slide, 0.45, y, 12.43, h, [s], fill=DGRAY, color=WHITE, bold=True, size=size, align=CENTER, name='まとめ')


def minimap(slide, current):
    """The map of the talk at the top right, with the current step in black."""
    x0, y0, w, h = 9.0, 1.2, 3.88, 0.6
    rect(slide, x0, y0, w, h, LGRAY, line=LINE, name='発表の流れ')
    text(slide, x0, y0 + 0.03, w, 0.2, ['本発表の流れ'], size=8, bold=True, color=DGRAY, align=CENTER, anchor=MIDDLE)
    bw = 0.64
    gap = (w - 0.12 - 5 * bw) / 4
    for i, label in enumerate(FLOW):
        bx = x0 + 0.06 + i * (bw + gap)
        now = label == current
        cell(slide, bx, y0 + 0.26, bw, 0.27, [label], fill=BLACK if now else SILVER, color=WHITE if now else MGRAY,
             bold=now, size=8.5, align=CENTER, pad=0.02)
        if i < 4:
            tri(slide, bx + bw + gap / 2 - 0.035, y0 + 0.36, 0.07, 0.07, GRAY, rot=90)


def frame(slide, n, title, subtitle, step=None, note=None):
    """Title (a sentence, given as its lines), the coloured subheading, the map, the footnote and the page number."""
    text(slide, 0.45, 0.28, 12.43, 0.9, list(title), size=22, bold=True, spacing=30, name='タイトル')
    text(slide, 0.45, 1.33, 8.3, 0.4, [subtitle], size=15, bold=True, color=ACCENT, anchor=MIDDLE, name='小見出し')
    if step:
        minimap(slide, step)
    if note:
        text(slide, 0.45, 6.88, 12.43, 0.25, [note], size=9, color=MGRAY, name='脚注')
    text(slide, 0.45, 7.14, 1.0, 0.22, [str(n)], size=9, color=MGRAY, name='ページ番号')
    text(slide, 6.88, 7.14, 6.0, 0.22, ['情報工学工房 2026　LLM-BLM'], size=9, color=MGRAY, align=RIGHT)


def style_axis_font(axis, size=11):
    set_font(axis.tick_labels.font, size, False, BLACK)
    axis.format.line.color.rgb = rgb(LINE)


def axis_title(axis, title):
    axis.has_title = True
    tf = axis.axis_title.text_frame
    tf.text = title
    set_font(tf.paragraphs[0].runs[0].font, 10.5, False, DGRAY)


def bar_chart(slide, x, y, w, h, cats, vals, colors, fmt='0', vmin=0, vmax=None, name=None):
    data = CategoryChartData()
    data.categories = cats
    data.add_series('値', vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data)
    if name:
        gf.name = name
    ch = gf.chart
    set_font(ch.font, 11, False, BLACK)
    ch.has_legend = False
    ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = 60
    plot.vary_by_categories = False
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = fmt
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    set_font(dl.font, 11.5, True, BLACK)
    ser = plot.series[0]
    ser.invert_if_negative = False
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
        pt.format.line.fill.background()
    # keep the colour of negative bars (PowerPoint draws them white unless each point says otherwise)
    for dpt in ser._element.findall(qn('c:dPt')):
        if dpt.find(qn('c:invertIfNegative')) is None:
            dpt.find(qn('c:idx')).addnext(dpt.makeelement(qn('c:invertIfNegative'), {'val': '0'}))
    va = ch.value_axis
    va.has_major_gridlines = False
    va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    va.visible = False
    ca = ch.category_axis
    style_axis_font(ca)
    ca.has_major_gridlines = False
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW     # labels below negative bars, not on them
    return gf


def line_chart(slide, x, y, w, h, cats, series, x_title=None, y_title=None, vmin=-3, vmax=3, major=1,
               fmt='+0"%";-0"%";0"%"', label_skip=None, name=None):
    data = CategoryChartData()
    data.categories = cats
    for label, vals, _ in series:
        data.add_series(label, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y), Inches(w), Inches(h), data)
    if name:
        gf.name = name
    ch = gf.chart
    set_font(ch.font, 11, False, BLACK)
    ch.has_title = False
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    set_font(ch.legend.font, 11, False, BLACK)
    for ser, (_, _, color) in zip(ch.plots[0].series, series):
        ser.smooth = False
        ser.marker.style = XL_MARKER_STYLE.NONE
        ser.format.line.color.rgb = rgb(color)
        ser.format.line.width = Pt(2.25)
    va = ch.value_axis
    va.minimum_scale = vmin
    va.maximum_scale = vmax
    va.major_unit = major
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb('E3E4E3')
    va.tick_labels.number_format = fmt
    va.tick_labels.number_format_is_linked = False
    style_axis_font(va)
    ca = ch.category_axis
    style_axis_font(ca)
    ca.has_major_gridlines = False
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    if label_skip:       # a label and a tick every label_skip categories (each label gets that much width)
        ax = ca._element
        for tag_name, val in (('c:tickLblSkip', label_skip), ('c:tickMarkSkip', label_skip)):
            el = ax.makeelement(qn(tag_name), {'val': str(val)})
            nm = ax.find(qn('c:noMultiLvlLbl'))
            if nm is not None:
                nm.addprevious(el)
            else:
                ax.append(el)
    if y_title:
        axis_title(va, y_title)
    if x_title:
        axis_title(ca, x_title)
    return gf


def blank(prs):
    return prs.slides.add_slide(next(l for l in prs.slide_layouts if l.name == 'Blank'))


# ---------- the deck ----------

def build():
    prs = Presentation(str(BASE))
    # keep the title slide, drop the rest
    ids = prs.slides._sldIdLst
    for sid in list(ids)[1:]:
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)
    prs.slides[0].notes_slide.notes_text_frame.text = (
        'AI に株を選ばせたら儲かるか、という研究の発表です。10 月 9 日から 27 日までに調べた結果を、7 分で話します。')

    # 2 ---------------------------------------------------------------- the whole study
    s = blank(prs)
    frame(s, 2, ('LLM は決算発表文から株価材料を読み取れるが、', '運用成績の向上に結び付く証拠は、検証を重ねるほど弱まった'),
          '研究の全体像')
    steps = ['① 数値予想を使う', '② 決算発表文を読む', '③ 大型株で運用', '④ 小型株で運用']
    x0, step, cw = 1.55, 2.86, 2.76
    for i, (label, fill) in enumerate(zip(steps, [GRAY, MGRAY, DGRAY, ACCENT])):
        x = x0 + i * step
        shape(s, MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, 1.95, cw if i == 3 else step + 0.1, 0.6, fill,
              name=f'段階{i + 1}')
        text(s, x + (0.12 if i == 0 else 0.32), 1.95, 2.4, 0.6, [label], size=13.5, bold=True, color=WHITE,
             align=CENTER, anchor=MIDDLE)
    label_block(s, 0.45, 2.72, 1.0, 1.25, '検証内容', size=12.5)
    label_block(s, 0.45, 4.07, 1.0, 1.45, '結果', size=12.5)
    what = ['直近 1 か月の日次リターンを渡し、翌月のリターンを数値で予想させて BLM のビューとする',
            '決算発表文を採点させ、発表を挟む 2 取引日の株価反応を説明できるかを検定',
            '点数を月次・日次の BLM に入れ、時価総額加重を上回るかを検定（S&P500）',
            '情報の織り込みが遅い小型株で、発表後の値動きの予測と BLM の超過収益を検定']
    found = [('直近の値動きをなぞるのみ', '予想と入力月の平均リターンの相関 0.89。固有の予測力は検出できず'),
             ('株価反応を有意に説明', '+120bp／1 標準偏差（t 9.4）。アナリスト予想などを統制しても残存'),
             ('市場平均を上回らず', '情報は発表当日に株価へ織り込まれ、翌日以降には残らない'),
             ('2 標本で優位、3 標本目で確認できず', '事前登録した検定は不支持。変動率の予測のみ 3 標本で成立')]
    for i in range(4):
        x = x0 + i * step
        cell(s, x, 2.72, cw, 1.25, [what[i]], size=11)
        head, detail = found[i]
        cell(s, x, 4.07, cw, 1.45, [[(head, {'bold': True, 'size': 12.5, 'color': ACCENT if i == 3 else BLACK})],
                                     [(detail, {'size': 11})]], line=ACCENT if i == 3 else None, after=3)
    funnel(s, 1.55, 5.64, 11.34, 0.28)
    closing(s, 6.02, 0.62, '読み取りは有効だが超過収益は立証できず、3 標本で成立したのはリスク（変動率）の予測のみ', size=15.5)
    s.notes_slide.notes_text_frame.text = (
        '最初に結論です。数字だけを見せると、AI は先月の動きをまねするだけでした。決算の文章を読ませると、'
        '発表直後の株価と合う点数を付けられました。でも大型株では、市場もその日のうちに読んでいて儲かりません。'
        '小型株では 2 つの標本で勝てましたが、まだ見ていない 3 つ目の標本では確かめられませんでした。'
        '最後まで残ったのは、値動きの荒さ、つまりリスクの予測です。')

    # 3 ---------------------------------------------------------------- the original paper
    s = blank(prs)
    frame(s, 3, ('元論文の手法を再現した結果、LLM の数値予想は直近の値動きをなぞるのみで、', '固有の予測力は確認されなかった'),
          '元論文（arXiv:2504.14345）の再現', step='元論文',
          note='*1: gpt-4o-mini、S&P500 構成銘柄 約 490 社 × 30 か月（回答 44 万個）。IC は予想と翌月リターンの相関')
    section(s, 0.45, 1.95, 4.75, 'BLM の仕組み')
    oval(s, 0.5, 2.42, 2.4, 0.78, SILVER, ['市場の均衡', [('時価総額の比率', {'size': 10, 'bold': False})]], BLACK)
    text(s, 1.5, 3.2, 0.4, 0.3, ['＋'], size=14, bold=True, align=CENTER, anchor=MIDDLE)
    oval(s, 0.5, 3.5, 2.4, 0.78, '75787A', ['LLM の見通し', [('平均＝ビュー、分散＝Ω', {'size': 10, 'bold': False})]],
         WHITE)
    tri(s, 3.0, 3.2, 0.28, 0.28, ACCENT, rot=90)
    cell(s, 3.4, 2.6, 1.8, 1.48, ['Black–Litterman', [('で混合', {'size': 11, 'bold': False})]], fill=BLACK,
         color=WHITE, bold=True, size=12, align=CENTER, pad=0.05)
    tri(s, 4.17, 4.14, 0.26, 0.2, ACCENT, rot=180)
    cell(s, 3.4, 4.4, 1.8, 0.6, ['最適化 → 投資比率'], bold=True, size=11, align=CENTER, pad=0.05)
    bullets(s, 0.45, 5.25, 4.75, 1.5, [(0, '確信度 Ω が小さいほど、比率は LLM の予想どおりに動く'),
                                      (0, '元論文は Llama・Gemma・Qwen の予想で、等金額・平均分散を上回ったと報告')],
            size=11.5)
    col_heads(s, 1.95, [(5.5, 1.35, '観点'), (6.95, 3.5, '再現結果'), (10.95, 1.93, '示唆')])
    rows = [('A', 'なぞり', '予想は入力した値動きの写し', '予想と入力月の平均リターンの相関 0.89',
             '数値の中で最も目立つ信号を写すのみ'),
            ('B', '固有の予測力', '入力で説明できない部分は当たらない', 'IC −0.004（95% 区間の上限 +0.027）',
             '数値の処理は統計モデルに任せるべき'),
            ('C', '確信度 Ω', 'Ω が過小で、見通しがビューに偏る', '事後期待リターンの 99.9% がビュー側',
             '市場の均衡がほぼ使われない'),
            ('D', '最適化', 'SLSQP が途中で停止', '「収束」と報告するが、入力の僅かな差で解が変わる',
             '成績が計算の作りに左右される')]
    for i, (t, label, head, detail, means) in enumerate(rows):
        y = 2.42 + i * 1.07
        label_block(s, 5.5, y, 1.35, 0.97, label, size=12.5)
        tag(s, 5.5, y, t)
        cell(s, 6.95, y, 3.5, 0.97, [[('✓ ' + head, {'bold': True, 'size': 12})], [(detail, {'size': 10.5})]])
        arrow_circle(s, 10.7, y + 0.485)
        cell(s, 10.95, y, 1.93, 0.97, [means], size=10.5)
    s.notes_slide.notes_text_frame.text = (
        'Black-Litterman モデルは、市場の平均的な見方から出発して、予想に自信がある分だけ比率をずらす方法です。'
        '元の論文は、AI に過去 1 か月の値動きを見せて来月を予想させ、これに入れると勝てたと報告しています。'
        '再現すると、AI の予想は先月の値動きと相関 0.89 で、ほぼ写しているだけでした。それを除いた AI ならではの当たりは'
        'ほぼ 0。しかも自信が強すぎて、市場の見方を 99.9% 捨てていました。')

    # 4 ---------------------------------------------------------------- earnings releases
    s = blank(prs)
    frame(s, 4, ('そこで問いを改め、決算発表文を LLM に採点させたところ、', '点数は発表直後 2 日間の株価反応を有意に説明した'),
          '決算発表文の読み取り（実験 2・3）', step='決算文',
          note='*1: 統制は数値の驚き（SUE・売上）、辞書・FinBERT・TF-IDF、発表前リターン、規模、業種・四半期の固定効果。'
               '5 分位の図は探索で、判定には用いていない')
    section(s, 0.45, 1.95, 5.3, '読み取りと評価の流れ')
    flow = [('入力', '決算発表文（EX-99.1）の先頭 15,000 字。S&P500、2023-11〜2026-08 の 5,374 本'),
            ('読み取り', 'gpt-4o-mini（学習なし、温度 0）が語調・今後の見通し・ガイダンスの変化を −2〜+2 で回答'),
            ('出力', '合成点数 ＝ 語調 ＋ 見通し ＋ ガイダンスの変化'),
            ('評価', '発表を挟む 2 取引日の超過リターン（同業種の指数比）に回帰')]
    for i, (label, body) in enumerate(flow):
        y = 2.42 + i * 0.9
        label_block(s, 0.45, y, 1.15, 0.8, label, size=12)
        cell(s, 1.68, y, 4.07, 0.8, [body], size=11)
    label_block(s, 0.45, 6.05, 1.15, 0.72, '事前登録', fill=ACCENT, size=12)
    cell(s, 1.68, 6.05, 4.07, 0.72, ['株価反応を計算する前に計画と評価コードを固定し、1 回のみ実行。以降の実験も同方式'],
         fill=WHITE, line=ACCENT, size=11)
    section(s, 6.1, 1.95, 6.78, '点数の 5 分位別：発表を挟む 2 日間の超過リターン')
    bar_chart(s, 6.1, 2.32, 6.78, 2.95, ['1（低）', '2', '3', '4', '5（高）'], QUINTILE,
              [DGRAY, DGRAY, DGRAY, ACCENT, ACCENT], fmt='+0.00"%";-0.00"%"', vmin=-2.5, vmax=2.0, name='5 分位')
    results = [('統制後の係数（実験 2）', '／1 標準偏差（t 9.4）', '+120bp'),
               ('＋アナリスト予想（実験 3）', '（t 9.4）、効果の 97% が残存', '+116bp')]
    for i, (label, rest, coef) in enumerate(results):
        y = 5.42 + i * 0.68
        label_block(s, 6.1, y, 2.6, 0.6, label, size=11.5)
        cell(s, 8.78, y, 4.1, 0.6, [[(coef, {'bold': True, 'size': 15, 'color': ACCENT}), (rest, {})]], size=12)
    s.notes_slide.notes_text_frame.text = (
        'そこで、AI が得意な文章を読ませることにしました。決算発表は会社が 3 か月ごとに出す成績表で、出た直後に株価が'
        '大きく動きます。gpt-4o-mini に読ませて採点させると、点数の高い発表ほど 2 日間の株価が上がり、上と下の 2 割で'
        '約 2.9% の差がありました。t 値は 9.4 で、偶然ではまず起きません。ズルをしないように、株価の反応を計算する前に'
        '計画とプログラムを固定して、1 回だけ実行しています。')

    # 5 ---------------------------------------------------------------- which model reads best
    s = blank(prs)
    frame(s, 5, ('7 つの手法を同一条件で比較した結果、学習なしの大規模 LLM が最も高い説明力を示し、', '先読みのない小型 LLM も従来手法を上回った'),
          '文章を数値化する手法の比較（実験 4・5・10・11）', step='決算文',
          note='*1: 説明力は、数値の驚き等の統制に加えたときの決定係数の増分（ポイント）。幅は実験ごとの値で、'
               '学習ありの手法はいずれも 2 日間の反応を正解とした')
    cols = [(0.95, 2.85, 'モデル'), (3.9, 2.85, '何を学習させたか（入力 → 正解）'), (6.85, 3.0, '学習の方法'),
            (9.95, 2.93, '説明力（R² の増分）')]
    col_heads(s, 1.92, cols)
    bar_x, scale = 10.0, 0.45                      # inches per point of R²
    for i, (model, model_sub, what, what_sub, how, how_sub, lo, hi, llm) in enumerate(MODELS):
        y = 2.38 + i * 0.49
        rect(s, 0.45, y, 12.43, 0.46, LGRAY, name=f'行{i + 1}')
        tag(s, 0.55, y + 0.08, str(i + 1), fill=ACCENT if llm else DGRAY, d=0.3, size=11)
        for (x, w, _), (main, sub) in zip(cols[:3], [(model, model_sub), (what, what_sub), (how, how_sub)]):
            text(s, x, y, w, 0.46, [[(main, {'size': 12, 'bold': True})], [(sub, {'size': 9.5, 'color': MGRAY})]],
                 anchor=MIDDLE, spacing=[15, 12.5])
        solid, light = (ACCENT, '9CC3F5') if llm else (DGRAY, LINE)
        rect(s, bar_x, y + 0.14, max(lo * scale, 0.03), 0.18, solid)
        if hi > lo:
            rect(s, bar_x + lo * scale, y + 0.14, (hi - lo) * scale, 0.18, light)
        label = f'{lo:.2f}〜{hi:.2f}' if hi > lo else f'{lo:.2f}'
        text(s, bar_x + max(hi * scale, 0.03) + 0.08, y, 1.1, 0.46, [label], size=11, bold=True, anchor=MIDDLE)
    label_block(s, 0.45, 5.88, 1.0, 0.9, '示唆', size=12.5)
    rect(s, 1.55, 5.88, 11.33, 0.9, LGRAY)
    bullets(s, 1.68, 5.95, 5.5, 0.8, [
        (0, [('大規模 LLM は学習なしで最強', {'bold': True}), ('（汎用の読解力が有効）', {})]),
        (0, [('株価反応で追加学習すると伸びる', {'bold': True}), ('（FinBERT 0.24 → 0.84〜1.01）', {})])],
        size=11, anchor=MIDDLE)
    bullets(s, 7.3, 5.95, 5.5, 0.8, [
        (0, [('LoRA はリッジに劣る', {'bold': True}), ('（約 1.4 万本に対し学習量 256 万は過大）', {})]),
        (0, [('先読みのない LLM も従来 NLP 3 手法の合計を上回る', {'bold': True})])],
        size=11, anchor=MIDDLE)
    s.notes_slide.notes_text_frame.text = (
        '情報工学として一番面白いのがこの比較です。同じ発表文を 7 つの方法で数にして、株価の動きをどれだけ説明できるかを'
        '比べました。学習させたモデルは、どれも発表直後 2 日の株価の動きを正解にしています。一番強かったのは、学習させて'
        'いない gpt-4o-mini でした。FinBERT は、全体を株価で追加学習させると 4 倍ほど伸びました。一方、ChronoGPT の中まで'
        ' LoRA で学習させると、中を動かさずにリッジ回帰を置くだけの方法に負けました。約 1 万 4 千本の発表に対して、'
        '学習する重みが 256 万個と多すぎたためと考えています。また、2015 年までの文章しか知らない ChronoGPT でも、'
        '従来の方法 3 つを合わせたものに勝ちました。答えを覚えていたのではない、という確認にもなっています。')

    # 6 ---------------------------------------------------------------- no profit in large caps
    s = blank(prs)
    frame(s, 6, ('一方、大型株では情報が発表当日に株価へ織り込まれるため、', '読み取りを用いた運用は市場平均を上回らない'),
          '大型株（S&P500）での運用（実験 7・8・9）', step='大型株',
          note='*1: 左図は点数の上位・中位・下位 20% の発表（各約 1,075 本）の平均。同業種の指数との差を発表 5 日前から累積')
    section(s, 0.45, 1.95, 6.9, '点数 3 分類別の累積超過リターン')
    cats = [('発表日' if d == 0 else (f'{-d} 日前' if d < 0 else f'{d} 日後') if d % 5 == 0 else '') for d in DAYS]
    line_chart(s, 0.45, 2.32, 6.9, 3.5, cats,
               [('点数 上位 20%', PATH_TOP, ACCENT), ('中位 20%', PATH_MID, GRAY), ('下位 20%', PATH_LOW, BLACK)],
               x_title='発表からの日数（取引日）', y_title='同業種との差（累積、%）', label_skip=5, name='前後のグラフ')
    text(s, 3.75, 3.02, 3.4, 0.3, ['発表当日に動き終わり、以降は横ばい'], size=11, bold=True)
    section(s, 7.65, 1.95, 5.23, '運用方法別の結果')
    rows = [('発表翌日に購入', None, 'ほぼ動かない', '翌日の超過リターン +3.6bp（t 0.93）'),
            ('月次 BLM', '約 10 年', '市場平均と同等', '費用控除後 +0.003%／月（t 0.04）'),
            ('日次 BLM', None, '費用で大きく劣後', '費用控除後 年 −20%（1 日の回転率 42%）')]
    for i, (label, sub, head, detail) in enumerate(rows):
        y = 2.42 + i * 1.12
        label_block(s, 7.65, y, 1.6, 1.0, label, sub=sub, size=12)
        cell(s, 9.33, y, 3.55, 1.0, [[('✓ ' + head, {'bold': True, 'size': 12.5})], [(detail, {'size': 10.5})]])
    funnel(s, 0.45, 5.98, 12.43, 0.24)
    closing(s, 6.26, 0.52, '見通しが価値を持つには未反映の情報が必要 → 情報の織り込みが遅い小型株で検証')
    s.notes_slide.notes_text_frame.text = (
        'では儲かるか。左のグラフは、AI の点数で発表を 3 グループに分け、発表の 5 日前から 20 日後まで、同じ業種の株と比べて'
        '何 % 上がったかを足し上げたものです。点数の高い発表は上がり、低い発表は下がりますが、動くのはほぼ'
        '発表の日だけで、そのあとは横ばいです。翌日に買ってもほぼ動かず、毎月組み替えて約 10 年運用しても市場の平均と'
        'ほぼ同じ、毎日組み替えると手数料で年マイナス 20% でした。AI が読めることは、市場もその日のうちに読んでいます。'
        'そこで、注目が少なく、情報がゆっくり株価に入る小型株に移りました。')

    # 7 ---------------------------------------------------------------- small caps
    s = blank(prs)
    frame(s, 7, ('情報の織り込みが遅い小型株では、読み取りが発表後の値動きを予測し、', 'BLM に組み込むと比較対象を上回った'),
          '小型株での予測と運用（実験 S1・S2・R1・G）', step='小型株',
          note='*1: 片道 0.3% の売買費用を控除（2017-01〜2026-09）。ありえない日次リターンを除いた後の計算であり、探索的位置付け')
    section(s, 0.45, 1.95, 5.35, '事前登録した予測の検定（t 値）')
    for x, label in [(2.45, '元の 1,000 社'), (4.15, '別の 1,000 社')]:
        cell(s, x, 2.38, 1.65, 0.38, [label], fill=DGRAY, color=WHITE, bold=True, size=11, align=CENTER, pad=0.03)
    preds = [('ドリフト', '発表後 2〜60 日の値動き', [('3.20', '+0.76%／SD'), ('2.22', '+0.54%／SD')]),
             ('変動率', '発表後 2〜60 日の荒さ', [('4.01', '+1.5%／SD'), ('3.01', '+1.06%／SD')])]
    for i, (label, sub, vals) in enumerate(preds):
        y = 2.82 + i * 0.86
        label_block(s, 0.45, y, 1.92, 0.78, label, sub=sub, size=13)
        for x, (t, per_sd) in zip((2.45, 4.15), vals):
            cell(s, x, y, 1.65, 0.78, [[('t ' + t, {'bold': True, 'size': 15, 'color': ACCENT})],
                                       [(per_sd, {'size': 10, 'color': MGRAY})]], align=CENTER, pad=0.03)
    section(s, 0.45, 4.62, 5.35, 'BLM の構成')
    bullets(s, 0.45, 5.02, 5.35, 1.8, [
        (0, [('標本：', {'bold': True}), ('米国小型株 1,000 社（浮動株時価総額 1,001〜3,000 位）、決算発表文 約 1.2 万本', {})]),
        (0, [('見通し：', {'bold': True}), ('ドリフトの点数（ChronoGPT＋リッジ、先読みなし）', {})]),
        (0, [('比較対象：', {'bold': True}), ('範囲の全社を時価総額に比例して保有', {})]),
        (0, [('制約：', {'bold': True}), ('比較対象からの乖離（TE）は年 3% 以内、四半期ごとに組み替え', {})])],
        size=11)
    section(s, 6.1, 1.95, 6.78, '資産の推移（元の 1,000 社、2016 年末＝1、費用控除後）')
    cats, blm, bench, act = asset_paths('orig')
    _, blm_r, bench_r, _ = asset_paths('rep')
    line_chart(s, 6.1, 2.32, 6.78, 4.45, cats,
               [('BLM（読み取りの見通し）', blm, ACCENT), ('比較対象（範囲の全社）', bench, DGRAY)],
               y_title='資産（2016 年末＝1）', vmin=0, vmax=5, major=1, fmt='0', label_skip=12, name='資産の推移')
    cell(s, 7.15, 2.45, 3.75, 1.02, [[(f'BLM {blm[-1]:.2f} 倍 ／ 比較対象 {bench[-1]:.2f} 倍', {'bold': True, 'size': 12.5})],
                                     f'差は年 +{act * 100:.1f}%（費用控除後）',
                                     f'別の 1,000 社：{blm_r[-1]:.2f} 倍 ／ {bench_r[-1]:.2f} 倍'],
         fill=WHITE, line=DGRAY, size=10.5, name='最後の値')
    s.notes_slide.notes_text_frame.text = (
        '米国の小型株 1,000 社、約 1 万 2 千本の決算発表で、先読みのない ChronoGPT とリッジの点数を使いました。'
        '事前登録した検定で、点数は発表後 2〜60 日の値動き、つまりドリフトと、値動きの荒さを予測しました。'
        'どちらも、重ならない別の 1,000 社で再現しています。そこで、この点数を BLM の見通しに入れ、'
        '四半期ごとに組み直しました。右のグラフが資産の推移です。2016 年末を 1 とすると、売買の費用を引いた後で、'
        'BLM は約 3.8 倍、範囲の全社を時価総額に比例して持つ相手は 3.0 倍でした。年にすると 2.4% の差です。')

    # 8 ---------------------------------------------------------------- is the edge from the reading
    s = blank(prs)
    frame(s, 8, ('2 標本での優位は偽の点数では再現されず、読み取りの中身に由来するが、', '売買費用に対しては脆弱である'),
          '優位の要因の確認（実験 G、探索）', step='検証',
          note='*1: 2017-01〜2026-09 の 2,449 日。結果は「元の 1,000 社・別の 1,000 社」の順。費用は BLM と比較対象の双方の売買に適用')
    col_heads(s, 1.95, [(0.45, 1.7, '観点'), (2.25, 3.05, '問い'), (5.4, 3.7, '結果（元・別の 1,000 社）'),
                        (9.6, 3.28, '解釈')])
    rows = [('A', '偽の点数', 'G1', '点数を銘柄間でシャッフルした偽の BLM（100 回）でも同等の成績となるか',
             '本物以上は 100 回中 0 回（p 0.01）', '偽物の平均は年 −1.0%・−1.3%、本物は +2.4%・+2.2%',
             '優位は点数の中身に由来', 'BLM の仕組みだけでは勝てない'),
            ('B', '既知の要因', 'G2', '規模・割安・収益性・投資・市場・モメンタムの 6 因子で説明できるか',
             'α 年 +1.9%・+1.6%', 't 2.12・1.65。モメンタムへの傾きが強い',
             '一部は因子で説明、残余あり', '差の約 2 割が因子による'),
            ('C', '売買費用', 'G5', '売買費用が増えても優位は残るか',
             '片道 1% で有意でなくなる', '0.5% で +1.9%・+1.7%、1% で +0.8%・+0.5%',
             '費用に脆弱', '四半期ごとに保有の約 85% を入れ替えるため')]
    for i, (t, label, g, q, rh, rd, ih, idt) in enumerate(rows):
        y = 2.42 + i * 1.17
        label_block(s, 0.45, y, 1.7, 1.07, label, sub=f'（{g}）', fill=BLACK, size=13)
        tag(s, 0.45, y, t)
        cell(s, 2.25, y, 3.05, 1.07, [q], size=11)
        cell(s, 5.4, y, 3.7, 1.07, [[(rh, {'bold': True, 'size': 12})], [(rd, {'size': 10.5})]])
        arrow_circle(s, 9.35, y + 0.535)
        cell(s, 9.6, y, 3.28, 1.07, [[('✓ ' + ih, {'bold': True, 'size': 12})], [(idt, {'size': 10.5})]])
    funnel(s, 0.45, 5.98, 12.43, 0.24)
    closing(s, 6.26, 0.52, '2 標本は結果を見ながら繰り返し分析した探索的結果 → 未使用の標本で事前登録検定を実施')
    s.notes_slide.notes_text_frame.text = (
        '前のページの勝ちが AI の読み取りのおかげかを、3 つ確かめました。1 つ目、点数を銘柄の間でシャッフルした偽物を'
        ' 100 回作ると、本物に届いたものは 1 回もありません。2 つ目、規模や割安さ、モメンタムといった、よく知られた要因を'
        '除いても年 1.6〜1.9% 残りました。3 つ目、ただし費用には弱く、片道 1% にすると有意でなくなります。'
        'この 2 つの標本は何度も調べたので、ここまでは探索です。')

    # 9 ---------------------------------------------------------------- G3: the preregistered third sample
    s = blank(prs)
    frame(s, 9, ('未使用の 400 社で事前登録した検定では BLM の優位は確認されず、', '3 標本すべてで成立したのは変動率の予測のみであった'),
          '事前登録による確認（実験 G3）', step='検証',
          note='*1: 変動率・ドリフトは点数 1 標準偏差あたりの係数。偽の点数との比較でも、400 社では本物以上が 100 回中 12 回（p 0.13）')
    label_block(s, 0.45, 1.95, 1.0, 1.0, '検定の設計', size=12)
    rect(s, 1.55, 1.95, 11.33, 1.0, LGRAY)
    bullets(s, 1.7, 1.99, 11.1, 0.92, [(0, '元・別のいずれの標本にも含まれない残り 400 社を使用'),
                                      (0, 'データ掃除の規則と売買費用を計算前に固定し、1 回のみ実行'),
                                      (0, '400 社では検出力が約 7 割にとどまることも事前に明記')],
            size=11, after=1, anchor=MIDDLE)
    cols = [(3.95, 2.93, '元の 1,000 社', DGRAY), (6.94, 2.93, '別の 1,000 社', DGRAY),
            (9.93, 2.95, '未使用の 400 社（事前登録）', BLACK)]
    for x, w, label, fill in cols:
        cell(s, x, 3.15, w, 0.5, [label], fill=fill, color=WHITE, bold=True, size=12, align=CENTER)
    marks = {'✓': ACCENT, '✗': RED, '–': MGRAY}
    rows = [('変動率の予測', 'S2：発表後 2〜60 日の荒さ',
             [('✓', '+1.5%（t 4.01）'), ('✓', '+1.06%（t 3.01）'), ('✓', '+2.04%（t 4.38）')]),
            ('ドリフトの予測', 'S1：発表後 2〜60 日の値動き',
             [('✓', '+0.76%（t 3.20）'), ('✓', '+0.54%（t 2.22）'), ('✗', '+0.24%（t 0.66）')]),
            ('BLM の優位', '費用控除後、年率',
             [('–', '+2.4%（探索）'), ('–', '+2.2%（探索）'), ('✗', '+0.74%（p 0.25）')])]
    for i, (label, sub, vals) in enumerate(rows):
        y = 3.72 + i * 0.76
        label_block(s, 0.45, y, 3.42, 0.7, label, sub=sub, size=13)
        for (x, w, _, _), (mark, val), last in zip(cols, vals, (False, False, True)):
            col = marks[mark]
            cell(s, x, y, w, 0.7, [[(mark + '  ', {'color': col, 'size': 15}),
                                    (val, {'color': col if mark != '✓' else BLACK})]],
                 bold=True, size=13, align=CENTER, line=BLACK if last else None)
    funnel(s, 0.45, 6.06, 12.43, 0.2)
    closing(s, 6.3, 0.5, 'BLM 不支持の主因は組み方ではなく、前段のドリフト予測がこの 400 社で弱かったこと', size=14.5)
    s.notes_slide.notes_text_frame.text = (
        'そこで、まだ一度も使っていない残りの 400 社で、規則を先に固定して 1 回だけ検定しました。400 社では検出力が'
        ' 7 割ほどしかないことも、先に書いておきました。結果、BLM の勝ちは年 0.74% で、有意ではありませんでした。'
        '表の真ん中の行を見ると、その前の段階のドリフトの予測も、この 400 社では弱かった。つまり組み方ではなく、'
        '予測そのものが標本によって変わります。一方、値動きの荒さの予測は、3 つの標本すべてで成り立ちました。')

    # 10 --------------------------------------------------------------- conclusion
    s = blank(prs)
    frame(s, 10, ('LLM の読み取りは株価材料として有効だが、BLM による超過収益は立証に至らず、', '確実な成果はリスクの予測に限られる'),
          '結論と今後の課題')
    col_heads(s, 1.95, [(0.45, 1.6, '観点'), (2.15, 5.0, '本研究で言えること'), (7.65, 5.23, '残る課題・今後の方針')])
    rows = [('A', '読み取り', '決算発表文から株価材料を読める',
             'アナリスト予想・追加学習モデル・先読みのない LLM と比べても有効',
             '記憶の完全な排除', '先読みの排除は ChronoGPT による傍証にとどまる'),
            ('B', '大型株', '情報は発表当日に織り込まれる', '月次・日次の BLM は費用控除後に市場平均を上回らない',
             '日中データでの検証', '発表直後の数分〜数時間の反応は未検証'),
            ('C', '小型株の予測', '変動率の予測は 3 標本で成立', 'ドリフトの予測は 3 標本中 2 標本で成立',
             'リスク管理への応用', '変動率の予測を保有量や TE の配分に活用'),
            ('D', '小型株の BLM', '事前登録した検定では優位を確認できず', '2 標本の年 +2.2〜2.4% は探索的結果',
             '前向きの検証', '2026 年 11 月以降の未見の決算発表で再検証')]
    for i, (t, label, head, detail, todo, todo_detail) in enumerate(rows):
        y = 2.42 + i * 0.97
        label_block(s, 0.45, y, 1.6, 0.88, label, size=12.5)
        tag(s, 0.45, y, t)
        cell(s, 2.15, y, 5.0, 0.88, [[('✓ ' + head, {'bold': True, 'size': 12})], [(detail, {'size': 10.5})]])
        arrow_circle(s, 7.4, y + 0.44)
        cell(s, 7.65, y, 5.23, 0.88, [[('✓ ' + todo, {'bold': True, 'size': 12})], [(todo_detail, {'size': 10.5})]])
    closing(s, 6.33, 0.47, '事前登録と別標本での確認により、期待した結果が弱まる過程も含めて検証結果を示した', size=14)
    s.notes_slide.notes_text_frame.text = (
        'まとめです。AI は決算の文章を読めます。ただし大型株では、その情報は発表の日に株価に入ります。小型株では、'
        '値動きの荒さの予測は 3 つの標本すべてで成り立ちましたが、BLM で勝てることは、事前登録した検定では確かめられません'
        'でした。学んだのは、事前登録すると、うまくいかなかったことも隠せない、ということです。今後は、確かに残った'
        'リスクの予測を運用の管理に使うこと、11 月以降のまだ誰も見ていない決算で確かめることを考えています。以上です。')

    prs.save(str(OUT))
    print('saved', OUT)


if __name__ == '__main__':
    build()
