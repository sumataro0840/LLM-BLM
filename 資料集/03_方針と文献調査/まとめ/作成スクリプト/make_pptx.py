"""LLM-BLM_文献調査と設計方針.pptx を作る (内容は同じ名前の .tex と同じ)。

    LLM/.venv/Scripts/python.exe 作成スクリプト/make_pptx.py   (まとめ/ の直下で実行する)
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parent.parent / "LLM-BLM_文献調査と設計方針.pptx"
FONT = "BIZ UDPGothic"
INK = RGBColor(0x26, 0x25, 0x23)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x88, 0x84)
LINE = RGBColor(0xD9, 0xD6, 0xD0)
C = {  # PDF と同じ色
    "blue": (RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xE4, 0xEE, 0xF9)),
    "orange": (RGBColor(0xEB, 0x68, 0x34), RGBColor(0xFB, 0xEE, 0xDF)),
    "green": (RGBColor(0x1B, 0xAF, 0x7A), RGBColor(0xE3, 0xF3, 0xE9)),
    "yellow": (RGBColor(0xED, 0xA1, 0x00), RGBColor(0xFD, 0xF3, 0xDA)),
    "purple": (RGBColor(0x62, 0x50, 0xD6), RGBColor(0xEE, 0xE8, 0xFA)),
}
W, H = Emu(12192000), Emu(6858000)  # 16:9
MARGIN = Emu(560000)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
page = [0]


def font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def text_box(slide, x, y, w, h, paragraphs, size=16, color=INK, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paragraphs: 文字列か (文字列, {"level", "bold", "size", "color", "bullet"}) のリスト。"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    for i, p in enumerate(paragraphs):
        text, opt = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        level = opt.get("level", 0)
        bullet = opt.get("bullet", True if level or opt.get("b") else False)
        prefix = ("・" if level == 0 else "– ") if bullet else ""
        if level:
            para.level = level
        para.space_after = Pt(opt.get("after", 6))
        run = para.add_run()
        run.text = ("　" * level) + prefix + text
        font(run, opt.get("size", size - 2 * level), opt.get("bold", False), opt.get("color", color))
    return box


def new_slide(title, kicker=None):
    slide = prs.slides.add_slide(BLANK)
    page[0] += 1
    if kicker:
        text_box(slide, MARGIN, Emu(330000), W - 2 * MARGIN, Emu(300000), [(kicker, {"size": 12, "color": C["blue"][0], "bold": True})])
    text_box(slide, MARGIN, Emu(560000), W - 2 * MARGIN, Emu(600000), [(title, {"size": 28, "bold": True})])
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, MARGIN, Emu(1200000), W - MARGIN, Emu(1200000))
    line.line.color.rgb = LINE
    line.line.width = Pt(1)
    text_box(slide, W - MARGIN - Emu(1500000), H - Emu(420000), Emu(1500000), Emu(260000),
             [(f"{page[0]}", {"size": 10, "color": MUTED})], align=PP_ALIGN.RIGHT)
    text_box(slide, MARGIN, H - Emu(420000), Emu(6000000), Emu(260000),
             [("LLM-BLM　文献調査と設計方針　2026-10-17", {"size": 10, "color": MUTED})])
    return slide


def card(slide, x, y, w, h, color, title, body, title_size=15, body_size=13):
    fg, bg = C[color]
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.color.rgb = fg
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Emu(140000)
    tf.margin_top = Emu(110000)
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    font(r, title_size, True, fg)
    p.space_after = Pt(4)
    for line in body if isinstance(body, list) else [body]:
        q = tf.add_paragraph()
        q.space_after = Pt(3)
        r = q.add_run()
        r.text = line
        font(r, body_size, False, INK)
    return shape


def node(slide, x, y, w, h, color, text, size=12):
    fg, bg = C[color]
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.1
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg
    shape.line.color.rgb = fg
    shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(60000)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        font(r, size, i == 0, INK if i else fg)
    return shape


def arrow(slide, a, b, start="right", end="left"):
    def point(s, side):
        if side == "right":
            return s.left + s.width, s.top + s.height // 2
        if side == "left":
            return s.left, s.top + s.height // 2
        if side == "bottom":
            return s.left + s.width // 2, s.top + s.height
        return s.left + s.width // 2, s.top
    x1, y1 = point(a, start)
    x2, y2 = point(b, end)
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = INK2
    c.line.width = Pt(1.25)
    tail = c.line._get_or_add_ln().makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    c.line._get_or_add_ln().append(tail)
    return c


def table(slide, x, y, w, rows, widths, size=12, header_color="blue", row_h=Emu(380000)):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, row_h * len(rows))
    tbl = shape.table
    total = sum(widths)
    for j, wd in enumerate(widths):
        tbl.columns[j].width = Emu(int(w * wd / total))
    fg, bg = C[header_color]
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Emu(80000)
            cell.margin_top = cell.margin_bottom = Emu(40000)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = fg if i == 0 else (RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF6, 0xF5, 0xF2))
            tf = cell.text_frame
            tf.word_wrap = True
            for k, line in enumerate(str(value).split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                r = p.add_run()
                r.text = line
                font(r, size, i == 0 or (j == 0), RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK)
    return shape


# ---------------------------------------------------------------- 1 表紙
s = prs.slides.add_slide(BLANK)
page[0] += 1
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Emu(260000), H)
band.fill.solid()
band.fill.fore_color.rgb = C["blue"][0]
band.line.fill.background()
text_box(s, Emu(900000), Emu(1900000), Emu(10400000), Emu(500000), [("情報工学工房　LLM-BLM", {"size": 16, "color": C["blue"][0], "bold": True})])
text_box(s, Emu(900000), Emu(2400000), Emu(10400000), Emu(1200000), [("文献調査と次の設計方針", {"size": 44, "bold": True})])
text_box(s, Emu(900000), Emu(3500000), Emu(10400000), Emu(900000),
         [("7 つの仮説の結果を踏まえ、何を問い、どう確かめるか", {"size": 20, "color": INK2})])
text_box(s, Emu(900000), Emu(5400000), Emu(10400000), Emu(500000), [("2026 年 10 月 17 日", {"size": 14, "color": MUTED})])

# ---------------------------------------------------------------- 2 要点
s = new_slide("この資料の要点")
items = [
    ("blue", "これまで", "LLM に翌月のリターンを数値で答えさせる構成は、統計を超える価値を示せなかった（予想の 81% は渡した月の写し）"),
    ("green", "新しい問い", "LLM が決算発表文から取り出した判断は、数値と LLM を使わない文章の分析を超える情報を持つか"),
    ("orange", "文献で変えること", "主な目的は「発表への反応」に。プロンプトに目的を書かない。2015〜2023 年は ChronoGPT で読む。取り下げ論文の引用を直す"),
    ("purple", "確かめ方", "手順を凍結した前向きの検証（2026 年 10 月以降の決算発表）を主な検定に。BLM は通過した信号を安全に比率へ変える部分"),
]
cw = (W - 2 * MARGIN - Emu(3 * 200000)) // 4
for i, (color, title, body) in enumerate(items):
    card(s, MARGIN + i * (cw + Emu(200000)), Emu(1600000), cw, Emu(3600000), color, title, body, 22, 19)

# ---------------------------------------------------------------- 3 これまでの結果
s = new_slide("これまでの結果：7 つの仮説", "背景")
rows = [
    ["仮説", "確かめたこと", "結果"],
    ["H01", "元論文の再現", "式は一致。BLM は LLM の予想を中央値 99.9% 採用。元論文の SLSQP は再現できない"],
    ["H02", "統計の基準との比較", "予想は渡した月の写し（決定係数 0.81）。固有の IC −0.004"],
    ["H03", "相対ビュー", "高ベータへの傾きは消えない（銘柄の間のなぞりが原因）"],
    ["H04", "Ω の較正", "BLM は設計どおり動くが、成績（年率 +0.53%）は並べ替えたビューと区別できない（p = 0.104）。ばらつきは大きさに使えず、順番の手がかりにはなる"],
    ["H05", "文章だけを渡す", "なぞりは消える（+0.77 → +0.05）。予測力は不十分"],
    ["H06", "決算発表文の点数", "5,429 本。IC +0.033（探索的）。数値を統制すると有意でない"],
    ["H07", "成績のフィードバック", "改善しない。一律に強気に寄った（日次 +0.12% → +0.29%）"],
]
table(s, MARGIN, Emu(1450000), W - 2 * MARGIN, rows, [1, 2.2, 8], 13, row_h=Emu(520000))
text_box(s, MARGIN, Emu(5850000), W - 2 * MARGIN, Emu(400000),
         [("評価は 2024 年 3 月〜2026 年 8 月の 30 か月（gpt-4o-mini の知識の期限より後）。IC 0.07 未満は検出できない。API 費用は計 6.7 ドル", {"size": 13, "color": INK2})])

# ---------------------------------------------------------------- 4 方法
s = new_slide("方法：3 者の方針と 4 者の文献調査", "方法")
a = node(s, MARGIN, Emu(2000000), Emu(2400000), Emu(1100000), "blue", "同じ依頼と資料\n見出しを揃える", 15)
b1 = node(s, Emu(3700000), Emu(1500000), Emu(2600000), Emu(900000), "green", "方針（3 者）\nClaude・GPT・Gemini", 14)
b2 = node(s, Emu(3700000), Emu(2700000), Emu(2600000), Emu(900000), "green", "文献調査（4 者）\n＋ Gemini Deep Research", 14)
c1 = node(s, Emu(7000000), Emu(2700000), Emu(2200000), Emu(900000), "orange", "照合\narXiv・OpenAlex・Crossref", 14)
d = node(s, Emu(9600000), Emu(2000000), Emu(2000000), Emu(1100000), "purple", "比較検討\n統合した方針", 15)
arrow(s, a, b1)
arrow(s, a, b2)
arrow(s, b2, c1)
arrow(s, b1, d)
arrow(s, c1, d)
text_box(s, MARGIN, Emu(4100000), W - 2 * MARGIN, Emu(2000000), [
    ("Claude は他の回答を読む前に自分の分を書き終えた（ただし実験の細部を知っている点で条件は同じでない）", {"b": True, "size": 15}),
    ("文献は arXiv の番号なら題名と著者まで突き合わせ、見つからないものは検索で個別に確かめた（LLM/refcheck.py）", {"b": True, "size": 15}),
    ("依頼の文書で H04 を逆に書く誤りがあった。GPT は添付のレポートから正しく読んだ。訂正を注記した", {"b": True, "size": 15, "color": C["orange"][0]}),
])

# ---------------------------------------------------------------- 5 照合の結果
s = new_slide("文献の照合：どのモデルにも誤りがあった", "方法")
rows = [
    ["", "文献", "検索", "照合で見つかった誤り"],
    ["Claude", "36", "約 35 回", "なし（H04 の読み違いを照合の段階で修正）"],
    ["GPT", "38", "14 回", "著者の取り違え 2 件。方針を変えた発見の多くはここから"],
    ["Gemini", "15", "8 語", "別の論文の番号、自分で「推定」とした DOI、見つからない文献、著者が Anonymous"],
    ["Deep Research", "17", "36 回", "著者の取り違え 4、題名の取り違え 2、主張の出どころの誤り、要旨にない数値、学術でない根拠"],
]
table(s, MARGIN, Emu(1500000), W - 2 * MARGIN, rows, [2, 1, 1.2, 8], 14, row_h=Emu(600000))
text_box(s, MARGIN, Emu(4800000), W - 2 * MARGIN, Emu(1300000), [
    ("よくある誤り：番号は実在するが別の論文／著者の取り違え／ある論文の主張を別の論文のものに／要旨にない数値", {"b": True, "size": 15}),
    ("照合で ok でも主張の中身は保証されない。この資料で使う主張は、要旨か本文で確かめたものに限った", {"b": True, "size": 15}),
])

# ---------------------------------------------------------------- 6 一致したこと
s = new_slide("4 者が一致したこと", "文献調査")
pts = [
    ("blue", "数値の予想はやめる", "LLM のリターン予想は直近を外挿しすぎ、楽観に寄る（Chen ほか 2024）。H02 の「写し」と同じ現象"),
    ("orange", "期限前は確認に使えない", "伏せても社名・日付が文脈から復元され、記憶と予測力を区別できない（Lopez-Lira ほか 2025）"),
    ("purple", "前向きの検証を主に", "先読みと研究者の選択の両方を防げる唯一の方法"),
    ("green", "非 LLM の 3 つの対照", "辞書（Loughran–McDonald）・TF-IDF の回帰・FinBERT。FinBERT は辞書を大きく上回る（Huang ほか 2023）"),
    ("yellow", "時刻の監査は必須", "同じ日の発表の相関で検定が過大に棄却する（Kolari・Pynnönen 2010）。悪いニュースは引け後に出やすい"),
]
cw = (W - 2 * MARGIN - Emu(2 * 200000)) // 3
for i, (color, title, body) in enumerate(pts):
    row, col = divmod(i, 3)
    card(s, MARGIN + col * (cw + Emu(200000)), Emu(1450000) + row * Emu(2350000), cw, Emu(2150000), color, title, body, 19, 16)

# ---------------------------------------------------------------- 7 方針を変える発見
s = new_slide("方針を変える 4 つの発見", "文献調査")
finds = [
    ("orange", "① 発表文は寄り付きで反映される",
     ["13.8 万本で、文章は利益の驚きと同等の情報。株価は寄り付きで内容を完全に反映（Wu ほか 2025）",
      "→ 主な目的変数を「発表への反応」に。ドリフトは副"]),
    ("purple", "② 目的を教えると測定が偏る",
     ["「株価の予測に使う」と伝えると語調の点数が目的に沿って偏る。改善は期限前だけ（Cao・Jiang・Xu 2026）",
      "→ プロンプトから株価への言及を除く"]),
    ("green", "③ 時点を区切った LLM が公開",
     ["ChronoGPT-Instruct：1999〜2024 年の年ごとの版、1.55B、MIT（He ほか 2025）",
      "→ 2015〜2023 年は匿名化でなくこれで読む"]),
    ("blue", "④ 引用していた論文が取り下げ",
     ["Kim・Muhn・Nikolaev（2024）は 2025 年 2 月に取り下げ（データの不整合）",
      "→ プロジェクトの資料 4 つの引用を直す"]),
]
cw = (W - 2 * MARGIN - Emu(200000)) // 2
for i, (color, title, body) in enumerate(finds):
    row, col = divmod(i, 2)
    card(s, MARGIN + col * (cw + Emu(200000)), Emu(1450000) + row * Emu(2400000), cw, Emu(2200000), color, title, body, 20, 17)

# ---------------------------------------------------------------- 8 分かれた点
s = new_slide("意見が分かれた点と、文献からの判断", "文献調査")
rows = [
    ["論点", "意見", "判断"],
    ["2015〜2023 年の文章", "前回は Gemini が匿名化で主に使う、GPT が使わない。今回は 4 者とも確認用にしない", "時点を区切ったモデルで読み「汚染のない下限」として報告。gpt-4o-mini の汚染は LAP の検定で診断"],
    ["意味のある最小の効果", "GPT：追加の rank IC 0.01＋経済効果。Gemini：0.03。Deep：0.05。Claude：検出力の実験で決める", "文献に共通の値はない。合成信号の検出力の実験で決め、GPT の基準を候補に"],
    ["BLM", "GPT は副次の評価に下げる。他は残す（Deep は Factor BLM）", "残す。役割は弱い信号での損と回転を抑えること。順位の戦略と同じ TE で比べる"],
    ["回答のばらつき", "GPT は二段階（大きさは誤差、比率は順位）が上限", "数値の予想をやめるので使わない。H04 は元論文の前提の検証として報告"],
]
table(s, MARGIN, Emu(1450000), W - 2 * MARGIN, rows, [2, 4.5, 4.5], 13, row_h=Emu(900000))

# ---------------------------------------------------------------- 9 研究の問い
s = new_slide("研究の問い", "設計方針")
card(s, MARGIN, Emu(1500000), W - 2 * MARGIN, Emu(1500000), "blue", "主な問い",
     "LLM が決算発表文から取り出した判断（語調・見通し・ガイダンスの変更）は、数値の驚き（SUE など）と、LLM を使わない文章の分析（辞書・TF-IDF・FinBERT）を超える情報を持つか", 18, 17)
text_box(s, MARGIN, Emu(3300000), W - 2 * MARGIN, Emu(2800000), [
    ("副次の問い", {"size": 17, "bold": True, "color": C["blue"][0]}),
    ("その情報は、寄り付きより後（取引できる時点）にも残るか", {"b": True, "size": 16}),
    ("先読みのない時点を区切ったモデルでも、同じ方向の情報が取れるか", {"b": True, "size": 16}),
    ("較正した BLM は、当たらない信号を自動で使わず、当たる信号だけを使うか（合成信号で）", {"b": True, "size": 16}),
    ("写し・ばらつき・フィードバックの結果を、元論文の前提の検証としてまとめる", {"b": True, "size": 16}),
])

# ---------------------------------------------------------------- 10 アーキテクチャ
s = new_slide("新しいアーキテクチャ", "設計方針")
y0 = Emu(1450000)
edgar = node(s, MARGIN, y0 + Emu(300000), Emu(2000000), Emu(800000), "blue", "EDGAR 8-K\nItem 2.02 EX-99.1", 13)
audit = node(s, MARGIN, y0 + Emu(1600000), Emu(2000000), Emu(1000000), "blue", "時刻の監査\n受理時刻 → ET\n寄り前／引け後", 13)
arrow(s, edgar, audit, "bottom", "top")
xs = Emu(3100000)
mids = []
for i, (color, text) in enumerate([
        ("green", "gpt-4o-mini（期限後）\n2023-11〜と前向きの記録"),
        ("green", "ChronoGPT（版＝前年末）\n2015〜2023"),
        ("yellow", "対照：辞書・TF-IDF・\nFinBERT"),
        ("yellow", "数値：SUE・売上の驚き・\n直前のリターン")]):
    n = node(s, xs, y0 + i * Emu(820000), Emu(3000000), Emu(700000), color, text, 12)
    arrow(s, audit, n)
    mids.append(n)
ev = node(s, Emu(7000000), y0 + Emu(600000), Emu(2400000), Emu(1700000), "purple",
          "イベントの評価\n(a) 発表への反応（主）\n(b) 寄り付き後のドリフト\n日付のクラスタ・置換", 13)
for n in mids:
    arrow(s, n, ev)
bl = node(s, Emu(9700000), y0 + Emu(600000), Emu(1950000), Emu(1700000), "orange",
          "通過した信号だけ\nBLM\nΩ∝Σ・較正・TE\nvs 順位の戦略", 13)
arrow(s, ev, bl)
text_box(s, MARGIN, Emu(4900000), W - 2 * MARGIN, Emu(1300000), [
    ("捨てる：数値の予想、ばらつきの Ω、予想役へのフィードバック、月次レポート、SLSQP、新しいモデルでの予想", {"b": True, "size": 14}),
    ("残す：Ω∝Σ と過去の当たり具合での較正、能動ベータ 0 と TE の上限、Clarabel、置換の対照、Holm の補正", {"b": True, "size": 14}),
    ("足す：時刻の監査、イベントの評価、非 LLM の 3 つの対照、ChronoGPT、LAP の検定、前向きの凍結、検出力の測定", {"b": True, "size": 14}),
])

# ---------------------------------------------------------------- 11 評価の枠組み
s = new_slide("評価の枠組み", "設計方針")
left = [
    ("目的変数", {"size": 17, "bold": True, "color": C["blue"][0]}),
    ("主：(a) 発表への反応（引け後なら翌日の寄り付きまで含む）", {"b": True, "size": 15}),
    ("副：(b) 寄り付き後 1・5・20 取引日の超過リターン（市場とセクターで調整）", {"b": True, "size": 15}),
    ("時刻", {"size": 17, "bold": True, "color": C["blue"][0]}),
    ("受理時刻を ET に。寄り前はその日、引け後は翌取引日が最初の取引日", {"b": True, "size": 15}),
    ("17 時 30 分 ET より後の受理は翌営業日扱いなので個別に確認", {"b": True, "size": 15}),
    ("統制と検定", {"size": 17, "bold": True, "color": C["blue"][0]}),
    ("SUE・売上の驚き・直前のリターン・規模・ベータ・セクターを統制", {"b": True, "size": 15}),
    ("日付ごとのクラスタリング、同じ日・同じ業種での置換、Holm の補正", {"b": True, "size": 15}),
]
text_box(s, MARGIN, Emu(1450000), Emu(6400000), Emu(4700000), left)
card(s, Emu(7300000), Emu(1450000), Emu(4330000), Emu(2150000), "purple", "前向きの検証（主な確認）",
     ["目的変数・点数・統制・除外の規則・窓・判定の基準・プロンプトとコードのハッシュを凍結",
      "2027 年 2 月まで結果を開けない"], 18, 16)
card(s, Emu(7300000), Emu(3800000), Emu(4330000), Emu(2150000), "green", "未来の情報の混入を防ぐ",
     ["期限前：時点を区切ったモデルだけ",
      "期限後：gpt-4o-mini",
      "最小の効果は合成信号の検出力の実験で決める"], 18, 16)

# ---------------------------------------------------------------- 12 ChronoGPT
s = new_slide("ChronoGPT の環境と最初の確認", "環境")
text_box(s, MARGIN, Emu(1450000), Emu(5400000), Emu(4600000), [
    ("この PC に環境を作り、2014〜2024 年末の 11 の版を取得（確認は 2022 年末の版。工房/LLM/chrono/）", {"b": True, "size": 15}),
    ("元の fp32（7.4GB）は 8GB の GPU に載らないので bf16（3.7GB）に変換。元のファイルの SHA-256 は公開値と一致", {"b": True, "size": 15}),
    ("点数付けは 1 本 0.1〜0.3 秒、GPU は最大 5.2GB", {"b": True, "size": 15}),
    ("2022 年の版でも大統領を「Bill Clinton」と答える（指示の学習データが 2000 年より前）→ 知識を問う使い方には向かない", {"b": True, "size": 15}),
    ("較正しないとガイダンスを 9 割以上「lower」と答える癖 → 中身のない文章で測った癖で割る「文脈による較正」を既定に", {"b": True, "size": 15}),
])
rows = [
    ["較正後（2022 年の版）", "良い発表の例", "悪い発表の例"],
    ["ガイダンス\nraise／lower／keep", "0.79／0.17／0.04", "0.33／0.54／0.13"],
    ["語調\npositive／negative", "0.79／0.20", "0.31／0.66"],
]
table(s, Emu(6300000), Emu(1500000), Emu(5330000), rows, [1.8, 1.8, 1.8], 13, "green", row_h=Emu(700000))
card(s, Emu(6300000), Emu(3900000), Emu(5330000), Emu(2000000), "orange", "研究での使い方",
     ["「据え置き」や中間の文章はまだ区別できない",
      "最後の層の埋め込み（1,536 次元）に、判断日より前のイベントだけで学習した回帰を置く（He ほか 2025 と同じ）",
      "ゼロショットの点数は補助"], 16, 13)

# ---------------------------------------------------------------- 13 次の実験
s = new_slide("次の実験（優先順）", "計画")
rows = [
    ["実験", "方法と成功の基準", "次の判断", "費用・時間"],
    ["P0 前向きの記録", "手順を凍結し、毎週 EDGAR の新しい発表文を取得・点数化。プロンプトから株価への言及を除く", "2027 年 2 月に一度だけ開ける", "数ドル・今週"],
    ["E1 時刻の監査", "5,429 件の受理時刻と取引日を確定。不明なものは主な分析から外す", "以後のすべての前提", "0 円・1 週"],
    ["E2 イベントの評価", "既存の点数で (a) 反応と (b) ドリフト。数値を統制して係数と区間", "(a) で残らなければ縮小", "0 円・1 週"],
    ["E3 非 LLM の対照", "辞書・TF-IDF・FinBERT を同じイベントに。LLM が 3 つを統制して残るか", "超えなければ安い方法へ", "0 円・1〜2 週"],
    ["E4 検出力と BLM", "2016〜2026 年に合成信号を入れ、必要な件数と較正した BLM の採用率を測る", "最小の効果と規模を決める", "0 円・1〜2 週"],
    ["E5 期限前の延長", "ChronoGPT で 2015〜2023 年を点数化。期限後で gpt-4o-mini と一致度。LAP の検定", "使えれば標本が約 4 倍", "0 円・2〜3 週"],
    ["E6 BLM への組み込み", "残った信号だけ。順位の戦略・等リスクと同じ TE・コストで比べる", "改善しなければ外す", "0 円・1〜2 週"],
]
table(s, MARGIN, Emu(1400000), W - 2 * MARGIN, rows, [2.1, 6, 2.6, 1.5], 12, row_h=Emu(580000))

# ---------------------------------------------------------------- 14 スケジュール
s = new_slide("スケジュール", "計画")
phases = [
    ("blue", "10 月後半", "P0 凍結と記録開始\nE1 時刻の監査\n引用の修正", "主な目的変数と基準を凍結したか"),
    ("green", "11 月", "E2・E3\n並行して E4", "文章の点数が数値と対照を超えるか"),
    ("yellow", "12 月", "E5 ChronoGPT で延長\nLAP の検定", "期限前の点数が期限後と同じ方向か"),
    ("orange", "1〜2 月", "E6（残った場合）\n前向きの記録を締める", "BLM を残すか"),
    ("purple", "2〜3 月", "前向きの検証を開ける\n最終報告", "凍結した基準で一度だけ判定"),
]
cw = (W - 2 * MARGIN - Emu(4 * 150000)) // 5
for i, (color, when, what, gate) in enumerate(phases):
    x = MARGIN + i * (cw + Emu(150000))
    card(s, x, Emu(1600000), cw, Emu(2000000), color, when, what.split("\n"), 20, 16)
    text_box(s, x, Emu(3850000), cw, Emu(1800000), [("判断点", {"size": 14, "bold": True, "color": C[color][0]}), (gate, {"size": 17})])

# ---------------------------------------------------------------- 15 やらないこと・リスク
s = new_slide("やらないこと・リスク", "計画")
card(s, MARGIN, Emu(1450000), Emu(5450000), Emu(3500000), "orange", "やらないこと", [
    "数値の予想のプロンプトの工夫・新しいモデルでの予想（外挿と先読みは工夫で消えない）",
    "予想役へのフィードバック（H07 で一律に強気に）",
    "成績を見て τ・TE・Ω を選ぶこと（過剰適合）",
    "2015〜2023 年を gpt-4o-mini と匿名化で読むこと（記憶を除けない）",
    "月次レポートの分析（米国の銘柄・セクターにほぼ触れない）"], 20, 17)
card(s, Emu(6180000), Emu(1450000), Emu(5450000), Emu(3500000), "purple", "リスクと未解決", [
    "前向きの標本は約 1,000 件。区間で報告し「有意でない」を「効果がない」と読まない",
    "ChronoGPT は文脈 1,792 トークンで読む力も弱い。見出しと冒頭の要約、段落ごとの点数で扱う",
    "寄り付きで反映されるなら取引できる価値はゼロに近い。それも結論になる",
    "独立に確かめられなかった文献（Liang 2026 など）は根拠に使わない"], 20, 17)

# ---------------------------------------------------------------- 16 決めること
s = new_slide("決めること", "次の一歩")
card(s, MARGIN, Emu(1600000), W - 2 * MARGIN, Emu(1500000), "purple", "1. 前向きの記録を今すぐ始めるか",
     "始めるなら、主な目的変数を「発表への反応」にするか（推奨：反応を主、ドリフトを副）。凍結の前にプロンプトから株価への言及を除く", 20, 16)
card(s, MARGIN, Emu(3400000), W - 2 * MARGIN, Emu(1500000), "blue", "2. 取り下げられた論文の引用を直すか",
     "LLM-BLM の資料 4 つ（総合報告書、19b 章、実験仮説、実験結果レポート）が Kim・Muhn・Nikolaev（2024）を成功例として引いている", 20, 16)

# ---------------------------------------------------------------- 17 参考文献
s = new_slide("主な参考文献", "付録")
refs = [
    "Cao, Jiang, Xu (2026). Seeing the Goal, Missing the Truth. NBER WP 35142; arXiv:2602.09504",
    "Chen, Green, Gulen, Zhou (2024). What Does ChatGPT Make of Historical Stock Returns? arXiv:2409.11540",
    "deHaan, Shevlin, Thornock (2015). Market (in)attention and the strategic scheduling and timing of earnings announcements. JAE 60",
    "Gao, Jiang, Yan (2025). Detecting Lookahead Bias in LLM Forecasts. arXiv:2512.23847",
    "He, Lv, Manela, Wu (2025). Chronologically Consistent LLMs. arXiv:2502.21206 ／ Instruction Tuning … arXiv:2510.11677",
    "Huang, Wang, Yang (2023). FinBERT. Contemporary Accounting Research 40(2)",
    "Kim, Muhn, Nikolaev (2024). Financial Statement Analysis with LLMs. arXiv:2407.17866（2025 年 2 月に取り下げ）",
    "Kolari, Pynnönen (2010). Event Study Testing with Cross-Sectional Correlation of Abnormal Returns. RFS 23(11)",
    "Lee, Kim, Kim, Lee (2025). LLM-Enhanced Black-Litterman Portfolio Optimization. CIKM 2025; arXiv:2504.14345",
    "Lopez-Lira, Tang, Zhu (2025). The Memorization Problem. arXiv:2504.14765",
    "Loughran, McDonald (2011). When Is a Liability Not a Liability? J. Finance 66(1)",
    "Siano (2025). The News in Earnings Announcement Disclosures. Management Science 71(11)",
    "Wu, Akin, Martineau, Grégoire, Veneris (2025). Extracting the Structure of Press Releases… arXiv:2509.24254",
    "Yu, Liu, Zhang, He (2026). Fast Numbers, Slow Language. arXiv:2606.29734",
]
text_box(s, MARGIN, Emu(1400000), W - 2 * MARGIN, Emu(4900000), [(r, {"size": 12, "after": 3}) for r in refs])

prs.save(OUT)
print(OUT, len(prs.slides), "slides")
